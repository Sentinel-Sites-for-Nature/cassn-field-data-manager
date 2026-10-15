"""Read the WI Processing Tracker's deployment-level Wildlife Insights status.

Program reporting (summary stats and the status dashboard) lives in the
separate cassn-reporting repository; this copy serves Box metadata maintenance
such as ``utils/backfill_wi_provenance.py``.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class WISubmissionTracker:
    """Deployment-level WI state loaded from the WI Processing Tracker."""

    source: Path
    status_by_deployment: dict[str, str]
    warnings: tuple[str, ...] = ()

    def submitted(self, deployment_id: str) -> bool | None:
        status = self.status_by_deployment.get(deployment_id)
        if status is None:
            return None
        normalized = status.strip().lower()
        if normalized.startswith("wi -"):
            return True
        if normalized == "box":
            return False
        return None


def _tracker_from_rows(source: Path, rows: list[dict[str, str]]) -> WISubmissionTracker:
    statuses: dict[str, str] = {}
    warnings: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        deployment_id = str(row.get("Deployment") or "").strip()
        status = str(row.get("Status") or "").strip()
        if not deployment_id:
            continue
        if deployment_id in statuses and statuses[deployment_id] != status:
            raise ValueError(f"{source}: conflicting WI statuses for {deployment_id!r}")
        statuses[deployment_id] = status
        normalized = status.lower()
        if normalized != "box" and not normalized.startswith("wi -"):
            warnings.append(f"{source.name} row {row_number}: unrecognized WI status {status!r}")
    if not statuses:
        raise ValueError(f"{source}: no Deployment rows found")
    return WISubmissionTracker(source, statuses, tuple(warnings))


def load_wi_submission_tracker(path: Path) -> WISubmissionTracker:
    """Load the Deployment Tracker tab from a CSV or exported XLSX workbook."""
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"WI tracker does not exist: {source}")
    if source.suffix.lower() == ".csv":
        with source.open(newline="", encoding="utf-8-sig") as stream:
            reader = csv.DictReader(stream)
            if reader.fieldnames is None:
                raise ValueError(f"{source}: CSV has no header")
            rows = [
                {str(key): str(value or "").strip() for key, value in row.items() if key is not None}
                for row in reader
            ]
        return _tracker_from_rows(source, rows)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("WI tracker must be a .csv or .xlsx file")

    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if "Deployment Tracker" not in workbook.sheetnames:
            raise ValueError(f"{source}: no 'Deployment Tracker' sheet")
        values = workbook["Deployment Tracker"].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        rows = [
            {headers[index]: str(value or "").strip() for index, value in enumerate(row)}
            for row in values
        ]
    finally:
        workbook.close()
    return _tracker_from_rows(source, rows)
